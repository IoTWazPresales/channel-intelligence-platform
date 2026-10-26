"""Independent pure XLSX serialization check: no database or spreadsheet execution."""
from io import BytesIO
import json
from pathlib import Path
import sys
import subprocess
from xml.etree import ElementTree as ET
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "apps/api"))
from openpyxl import load_workbook
from app.services.report_export import build_xlsx_bytes

labels = [
    '=WEBSERVICE("https://example.invalid/fictional")',
    "='[fictional.xlsx]Sheet1'!A1", "=cmd|' /C fictional'!A0",
    '#REF!', '#DIV/0!', '\r=1+1', '\n=1+1', ' =1+1',
    '<f>1+1</f>&"fictional"', '＝1+1', '−12.5', '=1+1\nsecond line',
]
observed = []
previous = {}
exec(compile(subprocess.check_output(['git', 'show', 'HEAD:apps/api/app/services/report_export.py'], cwd=ROOT).decode('utf-8'), '<HEAD report_export>', 'exec'), previous)
for label in labels:
    payload = build_xlsx_bytes(title=label, metric_key='fictional', grains=['customer'], value=0,
                              rows=[{label: label, 'negative': -12.5, 'integer': 7, 'boolean': False}],
                              data_vintage={'label': label}, invariants=[label])
    book = load_workbook(BytesIO(payload), data_only=False)
    for sheet, address in [('Cover', 'A1'), ('Cover', 'A10'), ('Data', 'A1'), ('Data', 'A2')]:
        cell = book[sheet][address]
        assert cell.data_type == 's', (label, sheet, address)
        if '\r' not in label:
            assert cell.value == label, (label, sheet, address)
    assert book['Cover']['A1'].font.bold and book['Data']['A1'].font.bold
    assert book['Data']['B2'].value == -12.5 and book['Data']['B2'].data_type == 'n'
    assert book['Data']['C2'].value == 7 and book['Data']['C2'].data_type == 'n'
    assert book['Data']['D2'].value is False and book['Data']['D2'].data_type == 'b'
    with ZipFile(BytesIO(payload)) as archive:
        assert not any('externalLinks' in name for name in archive.namelist())
        for name in archive.namelist():
            if name.endswith('.xml') or name.endswith('.rels'):
                tree = ET.fromstring(archive.read(name))
                for node in tree.iter():
                    assert node.tag.rsplit('}', 1)[-1] not in ('f', 'externalLink', 'externalReference'), name
                    assert node.attrib.get('TargetMode') != 'External', name
    original = previous['build_xlsx_bytes'](title=label, metric_key='fictional', grains=[], value=label, rows=None, data_vintage={})
    original_title = load_workbook(BytesIO(original))['Cover']['A1']
    if '\r' in label:
        assert original_title.value == book['Cover']['A1'].value
        assert original_title.data_type == 's'
    observed.append({'label': label, 'roundtrip_label': book['Cover']['A1'].value,
                     'exact_text': book['Cover']['A1'].value == label,
                     'baseline_same_text': original_title.value == book['Cover']['A1'].value,
                     'literal_type': True, 'numeric_boolean_types': True, 'no_formula_or_external_relationship': True})
print(json.dumps({'cases': len(observed), 'observations': observed}, ensure_ascii=True, indent=2))
