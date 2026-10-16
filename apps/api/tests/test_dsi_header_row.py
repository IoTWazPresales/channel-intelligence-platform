"""DSI reads the header row the file really has (blank leading row/column, banner forms, uploader override).

Finding on N-0111 (job 1669): a Pinnacle export with row 1 empty and column A empty mapped fine at the
mapping step (headers on row 2) but the validate pipeline re-read the file with header=0, saw only
``Unnamed: N`` columns, and erased the saved mapping -> ``missing_product_identifier_mapping``, 0 rows read.
"""

from __future__ import annotations

import io
from datetime import date, datetime

import pandas as pd
import pytest
from openpyxl import Workbook
from sqlalchemy import func, select

from app.models.ingestion import ImportJob
from app.services.imports.dsi_workbook import (
    DSI_HEADER_ROW_OVERRIDES_KEY,
    build_dsi_workbook_structure,
    describe_header_read,
    header_row_overrides_for_file,
    load_dsi_workbook_sheet_frames,
    no_header_row_message,
)

PINNACLE_HEADERS = [
    "Invoice Number",
    "Invoice Date",
    "Debtor Acc Code",
    "Debtor Name",
    "Sales Rep Code",
    "Sales Rep Name",
    "WRH",
    "Category",
    "Product Code",
    "Product Description",
    "QTY",
    "LOT ID",
]


def _pinnacle_xlsx(rows: list[tuple[str, date, float]], *, product: str = "SKU-ALPHA-01") -> bytes:
    """Sheet "Sales": row 1 empty, column A empty, headers B2:M2, data from row 3."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    for offset, header in enumerate(PINNACLE_HEADERS):
        ws.cell(row=2, column=2 + offset, value=header)
    for i, (invoice, when, qty) in enumerate(rows):
        r = 3 + i
        ws.cell(row=r, column=2, value=invoice)
        ws.cell(row=r, column=3, value=datetime(when.year, when.month, when.day))
        ws.cell(row=r, column=4, value="BNVEC001")
        ws.cell(row=r, column=5, value="Metro Market Group")
        ws.cell(row=r, column=6, value="BN292")
        ws.cell(row=r, column=7, value="BARNE MARAIS")
        ws.cell(row=r, column=8, value=1)
        ws.cell(row=r, column=9, value="ASUSNBVIVO")
        ws.cell(row=r, column=10, value=product)
        ws.cell(row=r, column=11, value="ASUS VIVOBOOK")
        ws.cell(row=r, column=12, value=qty)
        ws.cell(row=r, column=13, value="ASOP-NORMAL")
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _sample_rows() -> list[tuple[str, date, float]]:
    sales = [(f"INV{n:04d}", date(2026, 3, 1 + n), float(n + 1)) for n in range(8)]  # 1..8 -> 36 units
    credits = [(f"CRN{n:04d}", date(2026, 4, 1 + n), float(-(n + 1))) for n in range(3)]  # -1,-2,-3 -> 6 units
    return sales + credits


def test_pinnacle_shape_header_on_row_two_blank_column_a_dropped() -> None:
    raw = _pinnacle_xlsx(_sample_rows())
    frames = load_dsi_workbook_sheet_frames("pinnacle.xlsx", raw)
    assert len(frames) == 1
    sheet_name, df, header_row = frames[0]
    assert sheet_name == "Sales"
    assert header_row == 1  # 0-based -> "row 2"
    assert list(df.columns) == PINNACLE_HEADERS  # no "Unnamed: 0" from the blank column A
    assert len(df) == 11
    assert int((df["QTY"] < 0).sum()) == 3

    structure = build_dsi_workbook_structure("pinnacle.xlsx", raw)
    entry = structure["sheets"][0]
    assert entry["header_row"] == 1
    assert entry["header_row_user_set"] is False
    assert entry["columns"] == PINNACLE_HEADERS


def test_pinnacle_shape_csv_blank_first_line_and_column() -> None:
    lines = ["", "," + ",".join(PINNACLE_HEADERS), ",INV1,2026-03-01,A,B,C,D,1,E,SKU-ALPHA-01,Desc,2,LOT"]
    frames = load_dsi_workbook_sheet_frames("pinnacle.csv", ("\n".join(lines) + "\n").encode("utf-8"))
    _sn, df, header_row = frames[0]
    assert header_row == 1
    assert list(df.columns) == PINNACLE_HEADERS
    assert len(df) == 1


def test_headers_without_dsi_hint_words_use_the_fullest_early_row() -> None:
    """Row 1 blank, headers carry none of the DSI hint words: fall back to the fullest early row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for col, h in enumerate(["Alpha", "Bravo", "Charlie", "Delta"], start=2):
        ws.cell(row=3, column=col, value=h)
    for col, v in enumerate([1, 2, 3, 4], start=2):
        ws.cell(row=4, column=col, value=v)
    bio = io.BytesIO()
    wb.save(bio)
    _sn, df, header_row = load_dsi_workbook_sheet_frames("odd.xlsx", bio.getvalue())[0]
    assert header_row == 2
    assert list(df.columns) == ["Alpha", "Bravo", "Charlie", "Delta"]


def test_uploader_header_row_override_wins_over_detection() -> None:
    raw = _pinnacle_xlsx(_sample_rows())
    # Uploader says the headers are on row 3 (the first data row) -> that row becomes the header.
    _sn, df, header_row = load_dsi_workbook_sheet_frames(
        "pinnacle.xlsx", raw, header_row_overrides={"Sales": 2}
    )[0]
    assert header_row == 2
    assert "Invoice Number" not in list(df.columns)
    assert len(df) == 10
    structure = build_dsi_workbook_structure("pinnacle.xlsx", raw, header_row_overrides={"Sales": 2})
    assert structure["sheets"][0]["header_row_user_set"] is True


def test_override_lookup_is_per_file_and_sheet() -> None:
    job = ImportJob(
        template_slug="distributor_inventory",
        file_name="a.xlsx",
        staged_metadata={DSI_HEADER_ROW_OVERRIDES_KEY: {"a.xlsx::Sales": 4, "b.xlsx::Sales": 7, "a.xlsx::Stock": 1}},
    )
    assert header_row_overrides_for_file(job, "a.xlsx") == {"Sales": 4, "Stock": 1}
    assert header_row_overrides_for_file(job, "c.xlsx") == {}


def test_banner_layout_header_still_detected_on_row_20() -> None:
    """The existing form-block distributors: headers on row 20 (0-based 19) — unchanged by the blank-row work."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sell Out"
    ws["A5"] = "Subject: ASUS Weekly Sell Out Report"
    ws["A8"] = "Company Name"
    ws["B8"] = "MUSTEK-ZA-C"
    ws["A9"] = "Application Date"
    ws["B9"] = "2026W24"
    headers = [
        "Invoice Date",
        "Invoice No.",
        "Customer Code (Dealer Code)",
        "Dealer Name",
        "Model Name",
        "ASUS Part No.",
        "Quantity",
        "Unit Price",
    ]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=20, column=col, value=h)
    for col, v in enumerate(["2026-06-08", "1385729", "864966", "Hbits", "TP3407", "90NB14Y1", 1, 15304], start=1):
        ws.cell(row=21, column=col, value=v)
    bio = io.BytesIO()
    wb.save(bio)
    raw = bio.getvalue()

    _sn, df, header_row = load_dsi_workbook_sheet_frames("MUSTEK_SELLOUT WK24.xlsx", raw)[0]
    assert header_row == 19
    assert list(df.columns) == headers

    job = ImportJob(
        template_slug="distributor_inventory",
        file_name="MUSTEK_SELLOUT WK24.xlsx",
        staged_metadata={"dsi_workbook": build_dsi_workbook_structure("MUSTEK_SELLOUT WK24.xlsx", raw)},
    )
    assert "Headers read from row 20" in describe_header_read(job)


def test_describe_header_read_names_row_and_headers() -> None:
    raw = _pinnacle_xlsx(_sample_rows())
    job = ImportJob(
        template_slug="distributor_inventory",
        file_name="pinnacle.xlsx",
        staged_metadata={"dsi_workbook": build_dsi_workbook_structure("pinnacle.xlsx", raw)},
    )
    text = describe_header_read(job)
    assert "Headers read from row 2" in text
    assert "Invoice Number" in text and "Product Code" in text
    assert no_header_row_message(job) is None


def test_blank_sheet_says_no_header_row_found() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["B5"] = None
    bio = io.BytesIO()
    wb.save(bio)
    job = ImportJob(
        template_slug="distributor_inventory",
        file_name="blank.xlsx",
        staged_metadata={
            "dsi_workbook": {
                "sheets": [
                    {"sheet_name": "Sales", "sheet_key": "Sales", "columns": ["Unnamed: 0", "Unnamed: 1"], "header_row": 0}
                ]
            }
        },
    )
    msg = no_header_row_message(job)
    assert msg is not None
    assert "No header row found" in msg
    assert "Headers found on row" in msg


# --- cip_test end to end -------------------------------------------------------------------------


def _new_pinnacle_job(source_id: int, raw: bytes, filename: str) -> int:
    from app.db.session_sync import SessionLocal
    from app.models.ingestion import RawFileMetadata
    from app.storage.local import get_storage_backend

    with SessionLocal() as db:
        job = ImportJob(
            source_id=source_id,
            template_slug="distributor_inventory",
            import_mode="validate",
            status="pending",
            stage="uploaded",
            file_name=filename,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            staged_metadata={"dsi_workflow_mode_explicit": "historical"},
        )
        db.add(job)
        db.flush()
        key = f"imports/test/{job.id}/{filename}"
        get_storage_backend().save(key, raw, job.content_type)
        db.add(RawFileMetadata(job_id=job.id, storage_key=key, byte_size=len(raw), checksum=None))
        db.commit()
        return int(job.id)


PINNACLE_MAPPING = {
    "Invoice Number": "invoice_no",
    "Invoice Date": "transaction_date",
    "Product Code": "product_identifier",
    "QTY": "quantity_sold",
    "Debtor Name": "dealer_group_token",
}


def _confirm_dist01(db, job: ImportJob) -> None:
    from app.models.dimensions import DimDistributor
    from app.services.imports.dsi_file_distributor import set_dsi_file_distributors

    dist = db.scalar(select(DimDistributor).where(DimDistributor.code == "DIST-01"))
    assert dist is not None
    set_dsi_file_distributors(
        job,
        {
            job.file_name: {
                "token": "DIST-01",
                "reason": "steward_assigned",
                "confirmed": True,
                "distributor_id": int(dist.id),
                "distributor_name": dist.name,
            }
        },
    )


def _map_and_apply(job_id: int) -> ImportJob:
    from app.db.session_sync import SessionLocal
    from app.ingestion.pipeline import process_import_job_sync
    from app.services.imports.dsi_mapping_workflow import infer_dsi_job_sync

    with SessionLocal() as db:
        job = infer_dsi_job_sync(db, job_id)
        assert job.stage == "dsi_mapping_ready"
        # What the mapping step shows the uploader: the real headers, not "Unnamed: N".
        assert job.file_headers == PINNACLE_HEADERS
        assert job.staged_metadata["dsi_workbook"]["sheets"][0]["header_row"] == 1

        job.field_mapping = dict(PINNACLE_MAPPING)
        job.import_mode = "apply"
        _confirm_dist01(db, job)
        db.add(job)
        db.commit()
        return process_import_job_sync(db, job_id)


@pytest.fixture(scope="module")
def dsi_source_id() -> int:
    from tests.test_dsi_batch import _seed_dsi_source_for_batch_e2e

    return _seed_dsi_source_for_batch_e2e()


def _fact_totals() -> tuple[float, float, int, int]:
    """(units sold, units credited, sell-out rows, return rows) for the synthetic product/customer."""
    from app.db.session_sync import SessionLocal
    from app.models.dimensions import DimCustomer, DimProduct
    from app.models.facts import FactReturns, FactSalesSellout

    with SessionLocal() as db:
        product_id = db.scalar(select(DimProduct.id).where(DimProduct.sku == "SKU-ALPHA-01"))
        customer_id = db.scalar(select(DimCustomer.id).where(DimCustomer.code == "CUST-1001"))
        sell_scope = (
            FactSalesSellout.product_id == product_id,
            FactSalesSellout.customer_id == customer_id,
            FactSalesSellout.period_start >= date(2026, 3, 1),
            FactSalesSellout.period_start <= date(2026, 4, 30),
        )
        ret_scope = (
            FactReturns.product_id == product_id,
            FactReturns.customer_id == customer_id,
            FactReturns.transaction_date >= date(2026, 3, 1),
            FactReturns.transaction_date <= date(2026, 4, 30),
        )
        sold = db.scalar(select(func.coalesce(func.sum(FactSalesSellout.units), 0)).where(*sell_scope))
        credited = db.scalar(select(func.coalesce(func.sum(FactReturns.return_quantity), 0)).where(*ret_scope))
        n_sold = db.scalar(select(func.count()).select_from(FactSalesSellout).where(*sell_scope))
        n_credited = db.scalar(select(func.count()).select_from(FactReturns).where(*ret_scope))
        return float(sold or 0), float(credited or 0), int(n_sold or 0), int(n_credited or 0)


def test_pinnacle_shape_maps_loads_nets_credit_notes_and_reupload_changes_nothing(dsi_source_id: int) -> None:
    rows = _sample_rows()
    raw = _pinnacle_xlsx(rows)
    expected_sold = float(sum(q for _i, _d, q in rows if q > 0))
    expected_credited = float(-sum(q for _i, _d, q in rows if q < 0))

    first = _map_and_apply(_new_pinnacle_job(dsi_source_id, raw, "Asus Sales and Inventory - synthetic.xlsx"))
    assert first.status == "completed", first.error_summary

    sold, credited, n_sold, n_credited = _fact_totals()
    assert (n_sold, n_credited) == (8, 3)
    assert sold == expected_sold
    assert credited == expected_credited
    assert sold - credited == expected_sold - expected_credited  # net units

    # Re-upload the same file as a new job (Pinnacle selected): the facts do not change.
    second = _map_and_apply(_new_pinnacle_job(dsi_source_id, raw, "Asus Sales and Inventory - synthetic.xlsx"))
    assert second.status == "completed", second.error_summary
    assert _fact_totals() == (expected_sold, expected_credited, 8, 3)


def test_wrong_header_row_load_says_why_and_keeps_saved_mapping(dsi_source_id: int) -> None:
    """Uploader points at a row that is not the header: the load reports the row and headers read."""
    from app.db.session_sync import SessionLocal
    from app.ingestion.pipeline import process_import_job_sync
    from app.models.ingestion import ImportRowResult
    from app.services.imports.dsi_mapping_workflow import infer_dsi_job_sync, set_dsi_header_row_sync

    job_id = _new_pinnacle_job(dsi_source_id, _pinnacle_xlsx(_sample_rows()), "Asus wrong header row.xlsx")
    with SessionLocal() as db:
        infer_dsi_job_sync(db, job_id)
        job = set_dsi_header_row_sync(db, job_id, filename=None, sheet_key=None, header_row_number=5)
        assert job.staged_metadata[DSI_HEADER_ROW_OVERRIDES_KEY]["Asus wrong header row.xlsx::Sales"] == 4
        sheet = job.staged_metadata["dsi_workbook"]["sheets"][0]
        assert sheet["header_row"] == 4 and sheet["header_row_user_set"] is True

        job.field_mapping = dict(PINNACLE_MAPPING)  # headers of the *right* row: not found on row 5
        _confirm_dist01(db, job)
        db.add(job)
        db.commit()
        process_import_job_sync(db, job_id)
        messages = db.scalars(
            select(ImportRowResult.message).where(
                ImportRowResult.job_id == job_id, ImportRowResult.severity == "error"
            )
        ).all()
        text = " ".join(messages)
        assert "Headers read from row 5" in text
        assert "set by you" in text
        refreshed = db.get(ImportJob, job_id)
        assert refreshed is not None
        assert refreshed.field_mapping == PINNACLE_MAPPING  # not silently erased


def test_header_row_endpoint_validates_range(dsi_source_id: int) -> None:
    from app.db.session_sync import SessionLocal
    from app.services.imports.dsi_mapping_workflow import infer_dsi_job_sync, set_dsi_header_row_sync

    job_id = _new_pinnacle_job(dsi_source_id, _pinnacle_xlsx(_sample_rows()), "Asus header range.xlsx")
    with SessionLocal() as db:
        infer_dsi_job_sync(db, job_id)
        with pytest.raises(ValueError):
            set_dsi_header_row_sync(db, job_id, filename=None, sheet_key=None, header_row_number=0)
        with pytest.raises(ValueError):
            set_dsi_header_row_sync(db, job_id, filename="nope.xlsx", sheet_key=None, header_row_number=2)
        job = set_dsi_header_row_sync(db, job_id, filename=None, sheet_key=None, header_row_number=2)
        assert job.file_headers == PINNACLE_HEADERS
