"""E4: a saved question retains its execution semantics and explicit context."""
from contextlib import asynccontextmanager
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from tests.test_dashboard_v2_authenticated import signed_in  # noqa: F401
from tests.test_report_schedule_recipient_scope import scheduled  # noqa: F401


@pytest.mark.parametrize("own,effective", [(None, "missing"), (None, ["None"]), ("None", [None])])
def test_null_owned_filter_cannot_be_omitted_or_replaced_by_literal_none(own, effective):
    from app.services.dashboard_definitions import DefinitionRefused, check_definition, definition_for_promotion
    query = {"measures": ["dsi_sellout.units"], "breakdowns": ["bu"], "filters": {"bu": [own]}}
    definition = check_definition({"query": query}, tenant_id="default").definition
    execution = {"measures": query["measures"], "breakdowns": query["breakdowns"]}
    if effective != "missing":
        execution["filters"] = {"bu": effective}
    with pytest.raises(DefinitionRefused):
        definition_for_promotion(definition, tenant_id="default", execution=execution)


def test_explicit_null_owned_filter_is_preserved_by_server_promotion():
    from app.services.dashboard_definitions import check_definition, definition_for_promotion
    query = {"measures": ["dsi_sellout.units"], "breakdowns": ["bu"], "filters": {"bu": [None]}}
    definition = check_definition({"query": query}, tenant_id="default").definition
    assert definition_for_promotion(definition, tenant_id="default", execution=query).definition["query"] == query


def test_governed_default_cap_preserves_handler_order_and_does_not_mutate_cached_rows():
    from app.services.dashboard_definitions import check_definition
    from app.query.model.execute import present_handler_rows
    plan = check_definition({"query": {"measures": ["support_spend"], "breakdowns": ["customer"]}}, tenant_id="default").plan
    rows = [{"customer_id": i, "value": 501-i} for i in range(501)]
    presented, truncated = present_handler_rows(rows, plan)
    assert truncated and presented == rows[:500] and len(rows) == 501


@pytest.mark.anyio
async def test_implicit_model_order_and_default_cap_survive_the_legacy_handler(signed_in):
    from datetime import date
    from uuid import uuid4
    from app.api.deps import get_db
    from app.main import app
    from app.models.dimensions import DimCustomer, DimProduct
    from app.models.facts import FactSalesSellout
    s = signed_in
    factory = asynccontextmanager(app.dependency_overrides[get_db])
    async with factory() as db:
        tid = s.products[0].tenant_id
        customer = (await db.scalars(select(DimCustomer).where(DimCustomer.tenant_id == tid))).one()
        tag = uuid4().hex
        products = [DimProduct(tenant_id=tid, sku=f"e4-{tag}-{i}", name=f"E4 cap {i}", product_line=s.line_a) for i in range(500)]
        db.add_all(products)
        await db.flush()
        db.add_all([FactSalesSellout(tenant_id=tid, source_key=f"e4-{tag}-{i}", product_id=product.id,
                                  customer_id=customer.id, transaction_date=date(2026, 1, 5),
                                  period_start=date(2026, 1, 5), units=10000+i, revenue=0)
                    for i, product in enumerate(products)])
        await db.commit()  # Fictional quantities only; outer fixture rollback owns them.
    query = {"measures": ["dsi_sellout.units"], "breakdowns": ["product"]}
    expected = (await s.client.post("/api/v1/query/v2/execute", json=query, headers=s.headers["admin"])).json()
    assert expected["ok"] and expected["truncated"] and len(expected["rows"]) == 500
    response = await s.client.post("/api/v1/saved-reports", json={
        "name": "E4 implicit defaults", "visual": "table", "definition_v2": {"query": query},
    }, headers=s.headers["admin"])
    assert response.status_code == 200, response.text
    assert response.json()["definition_v2"]["v1_equivalent"]  # Retain the legacy handler authority.
    rid = response.json()["id"]
    response = await s.client.get(f"/api/v1/reports/saved/{rid}/export?format=xlsx", headers=s.headers["admin"])
    values = list(load_workbook(BytesIO(response.content))["Data"].values)
    assert len(values) == 501
    col = values[0].index("product_id")
    assert [row[col] for row in values[1:]] == [row["product_id"] for row in expected["rows"]]
    response = await s.client.post(f"/api/v1/reports/saved/{rid}/deliver", json={}, headers=s.headers["admin"])
    assert response.status_code == 200, response.text
    assert response.json()["data_vintage"]["returned_rows"] == 500
    assert response.json()["data_vintage"]["row_count"] == 502
    assert float(response.json()["value_preview"]) == expected["totals"]["dsi_sellout.units"]


@pytest.mark.parametrize("measure,breakdowns", [
    ("dsi_sellout.units", ["product"]), ("support_spend", ["customer"]),
])
@pytest.mark.parametrize("options", [
    {"limit": 1}, {"sort": {"dir": "asc"}}, {"sort": {"by": "PLACEHOLDER", "dir": "asc"}},
])
def test_explicit_sort_or_limit_never_uses_a_lossy_v1_mirror(measure, breakdowns, options):
    from app.services.dashboard_definitions import check_definition, runs_v1
    query = {"measures": [measure], "breakdowns": breakdowns, **options}
    if query.get("sort", {}).get("by"):
        query["sort"] = {"by": measure, "dir": "asc"}
    checked = check_definition({"query": query}, tenant_id="default")
    assert not checked.v1_equivalent
    assert not runs_v1(checked.definition)
    # Old rows already carry the incorrect server flag. Correct dispatch too.
    assert not runs_v1({**checked.definition, "v1_equivalent": True})


@pytest.mark.anyio
async def test_preexisting_false_equivalence_exports_and_delivers_the_same_limited_rows(signed_in, monkeypatch):
    from app.api.deps import get_db
    from app.api.v1.endpoints import report_exports
    from app.main import app
    from app.models.saved_reports import SavedReport
    s = signed_in
    query = {"measures": ["dsi_sellout.units"], "breakdowns": ["product"],
             "sort": {"by": "dsi_sellout.units", "dir": "asc"}, "limit": 1}
    expected = (await s.client.post("/api/v1/query/v2/execute", json=query, headers=s.headers["admin"])).json()
    assert expected["ok"] and expected["truncated"] and expected["totals"]["dsi_sellout.units"] == 1012
    response = await s.client.post("/api/v1/saved-reports", json={
        "name": "E4 limited fictional", "visual": "table", "definition_v2": {"query": query},
    }, headers=s.headers["admin"])
    assert response.status_code == 200, response.text
    rid = response.json()["id"]
    factory = asynccontextmanager(app.dependency_overrides[get_db])
    async with factory() as db:
        row = await db.get(SavedReport, rid)
        row.definition_v2 = {**row.definition_v2, "v1_equivalent": True}
        row.metric_key, row.grains = "sellout_units", ["product"]
        await db.commit()  # Simulate a pre-existing persisted incorrect mirror.
    response = await s.client.get(f"/api/v1/reports/saved/{rid}/export?format=xlsx", headers=s.headers["admin"])
    assert response.status_code == 200, response.text
    wb = load_workbook(BytesIO(response.content))
    rows = list(wb["Data"].values)
    assert len(rows) == 2  # header plus the one displayed row, not the full v1 result
    assert rows[1][rows[0].index("dsi_sellout.units")] == 99
    assert "1012" in wb["Cover"]["A4"].value
    captured = {}
    real_pdf = report_exports.build_pdf_bytes
    def observed_pdf(**kwargs):
        captured.update(kwargs)
        return real_pdf(**kwargs)
    monkeypatch.setattr(report_exports, "build_pdf_bytes", observed_pdf)
    response = await s.client.get(f"/api/v1/reports/saved/{rid}/export?format=pdf", headers=s.headers["admin"])
    assert response.status_code == 200 and response.content.startswith(b"%PDF")
    assert captured["rows"] == expected["rows"] and captured["value"] == 1012
    response = await s.client.post(f"/api/v1/reports/saved/{rid}/deliver", json={}, headers=s.headers["admin"])
    assert response.status_code == 200, response.text
    assert float(response.json()["value_preview"]) == 1012
    assert response.json()["data_vintage"]["returned_rows"] == 1
    assert response.json()["data_vintage"]["truncated"] is True


@pytest.mark.anyio
async def test_legacy_saved_and_ad_hoc_export_and_delivery_forward_period_grain(signed_in, monkeypatch):
    from app.api.v1.endpoints import report_exports
    from app.services import report_delivery
    s = signed_in
    body = {"name": "E4 legacy month", "metric": "sellout_units", "grains": ["period"],
            "period_grain": "month", "filters": {}, "visual": "table"}
    response = await s.client.post("/api/v1/saved-reports", json=body, headers=s.headers["admin"])
    assert response.status_code == 200, response.text
    rid = response.json()["id"]
    seen = []
    real_query = report_exports.execute_query
    async def observed_query(*args, **kwargs):
        seen.append(kwargs.get("period_grain"))
        return await real_query(*args, **kwargs)
    monkeypatch.setattr(report_exports, "execute_query", observed_query)
    monkeypatch.setattr(report_delivery, "execute_query", observed_query)
    for url, method, payload in [
        (f"/api/v1/reports/saved/{rid}/export?format=xlsx", "GET", None),
        ("/api/v1/reports/export", "POST", {"metric": "sellout_units", "grains": ["period"], "period_grain": "month"}),
    ]:
        response = await s.client.request(method, url, json=payload, headers=s.headers["admin"])
        assert response.status_code == 200, response.text
        values = list(load_workbook(BytesIO(response.content))["Data"].values)
        assert len(values) == 3 and {"2026-01", "2026-02"} <= {str(cell) for row in values[1:] for cell in row}
    response = await s.client.post(f"/api/v1/reports/saved/{rid}/deliver", json={}, headers=s.headers["admin"])
    assert response.status_code == 200, response.text
    assert seen == ["month"] * 3


@pytest.mark.anyio
async def test_promotion_preserves_actual_window_filters_order_and_limit(signed_in):
    s = signed_in
    headers = s.headers["admin"]
    response = await s.client.post("/api/v1/dashboards", json={
        "name": "E4 snapshot", "filters_json": {"date": {"preset": "last_13_weeks"}},
    }, headers=headers)
    assert response.status_code == 200, response.text
    base = f"/api/v1/dashboards/{response.json()['id']}"
    query = {"measures": ["dsi_sellout.units"], "breakdowns": ["product"],
             "sort": {"by": "dsi_sellout.units", "dir": "asc"}, "limit": 1,
             "filters": {"product": [s.products[0].id, s.products[1].id]}}
    response = await s.client.post(f"{base}/widgets", json={
        "title": "E4 selected January", "visual": "table", "definition_v2": {"query": query},
    }, headers=headers)
    assert response.status_code == 200, response.text
    wid = response.json()["widgets"][0]["id"]
    # Unsaved/current card context differs from the persisted relative dashboard.
    execution = {**query, "filters": {"product": [s.products[0].id]},
                 "period": {"from": "2026-01-01", "to": "2026-01-31"}}
    response = await s.client.post(f"{base}/widgets/{wid}/promote", json={"execution": execution}, headers=headers)
    assert response.status_code == 200, response.text
    saved = response.json()["saved_report"]
    definition = saved["definition_v2"]
    assert definition["query"]["filters"] == execution["filters"]
    assert definition["dashboard"]["fixed_window"] == execution["period"]
    assert definition["query"]["sort"] == query["sort"] and definition["query"]["limit"] == 1
    assert not set(definition["query"].get("period", {})) & {"from", "to"}
    response = await s.client.post(f"/api/v1/reports/saved/{saved['id']}/deliver", json={}, headers=headers)
    assert response.status_code == 200 and float(response.json()["value_preview"]) == 13
    assert response.json()["data_vintage"]["period"] == execution["period"]
    # The widget retains its reusable original question and own restriction.
    widget = (await s.client.get(base, headers=headers)).json()["widgets"][0]
    assert widget["definition_v2"]["query"]["filters"] == query["filters"]
    assert "dashboard" not in widget["definition_v2"]


@pytest.mark.anyio
@pytest.mark.parametrize("invalid", ["missing", "stale", "widened", "partial_dates"])
async def test_promotion_refuses_missing_or_misleading_execution_context(signed_in, invalid):
    s = signed_in
    headers = s.headers["admin"]
    response = await s.client.post("/api/v1/dashboards", json={"name": "E4 refusal"}, headers=headers)
    base = f"/api/v1/dashboards/{response.json()['id']}"
    query = {"measures": ["dsi_sellout.units_per_week"], "filters": {"product": [s.products[0].id]}}
    response = await s.client.post(f"{base}/widgets", json={
        "title": "E4 dated rate", "definition_v2": {"query": query},
    }, headers=headers)
    assert response.status_code == 200, response.text
    wid = response.json()["widgets"][0]["id"]
    execution = {**query, "period": {"from": "2026-01-01", "to": "2026-01-31"}}
    if invalid == "stale":
        execution["measures"] = ["dsi_sellout.units"]
    elif invalid == "widened":
        execution["filters"] = {"product": [s.products[1].id]}
    elif invalid == "partial_dates":
        execution["period"] = {"from": "2026-01-01"}
    response = await s.client.post(f"{base}/widgets/{wid}/promote",
                                   json={} if invalid == "missing" else {"execution": execution}, headers=headers)
    assert response.status_code == 400, response.text
    assert not (await s.client.get("/api/v1/saved-reports", headers=headers)).json()["items"]


@pytest.mark.anyio
async def test_import_trigger_bypasses_cache_after_sorted_governed_report_routes_to_v2(scheduled):
    from app.models.cpor import CporCase, CporCaseLine
    from app.models.dimensions import DimCustomer
    from app.services.dashboard_definitions import check_definition
    from tests.test_report_schedule_recipient_scope import schedule, deliveries
    s = scheduled
    report = s.reports[0]
    async with s.factory() as db:
        case = (await db.scalars(select(CporCase).where(CporCase.tenant_id == report.tenant_id).order_by(CporCase.id.desc()))).first()
        customer = DimCustomer(tenant_id=report.tenant_id, code="E4-fictional", name="E4 second customer")
        db.add(customer)
        await db.flush()
        case.customer_id = customer.id
        checked = check_definition({"query": {"measures": ["support_spend"], "breakdowns": ["customer"],
                                             "sort": {"by": "support_spend", "dir": "asc"}, "limit": 1}}, tenant_id=report.tenant_id)
        stored = await db.get(type(report), report.id)
        stored.definition_v2, stored.metric_key = checked.definition, checked.v1.metric_key
        await db.commit()
    sid = await schedule(s, report, [s.users["admin"]])
    assert (await s.runner._deliver_schedule_async(sid, trigger="schedule", advance_clock=False))["ok"]
    first = (await deliveries(s, report))[-1]
    assert first.data_vintage["returned_rows"] == 1 and first.data_vintage["truncated"] is True
    assert float(first.value_preview) == 70
    async with s.factory() as db:
        line = (await db.scalars(select(CporCaseLine).where(CporCaseLine.case_id == case.id))).one()
        line.ttl_support_usd = 90
        await db.commit()
    assert (await s.runner._deliver_schedule_async(sid, trigger="import_event", advance_clock=False))["ok"]
    assert float((await deliveries(s, report))[-1].value_preview) == 100
