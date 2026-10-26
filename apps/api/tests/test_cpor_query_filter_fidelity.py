"""E1: governed CPOR filters constrain the source cohort, independently of grain.

Real PostgreSQL on cip_test, fictional records, one rolled-back transaction. No
commercial computations are mocked. No production connection or delivery send.
"""
from datetime import date
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.core.product_line_scope import request_product_lines, request_role
from app.core.tenant_scope import request_tenant_id
from app.db.session import AsyncSessionLocal
from app.models.cpor import CporCase, CporCaseLine
from app.models.dimensions import DimCustomer, DimProduct
from app.models.saved_reports import SavedReport
from app.query.cache import clear_query_cache
from app.query.engine import execute_query
from app.query.model.execute import run_v2
from app.query.model.graph import V2Request
from app.semantics.registry import load_catalog
from app.services.report_delivery import _v2_payload, deliver_saved_report
from app.services.report_export import build_xlsx_bytes


@pytest.mark.anyio
async def test_filters_and_totals_at_every_advertised_cpor_grain(monkeypatch):
    tag = uuid4().hex[:10]
    tenant = f"e1-{tag}"
    async with AsyncSessionLocal() as db:
        assert (await db.execute(text("select current_database()"))).scalar_one() == "cip_test"
        await db.execute(text("INSERT INTO tenant (id, name) VALUES (:id, :name)"), {"id": tenant, "name": "E1 fictional tenant"})
        customers = [DimCustomer(tenant_id=tenant, code=f"E1-{tag}-{i}", name=f"Fictional {i}") for i in range(2)]
        products = [DimProduct(tenant_id=tenant, sku=f"E1-{tag}-{i}", name=f"Fictional {i}", product_line=bu) for i, bu in enumerate(("NB", "NR"))]
        db.add_all(customers + products)
        await db.flush()
        cases = []
        for i, customer in enumerate(customers):
            case = CporCase(tenant_id=tenant, case_code=f"E1-{tag}-{i}", customer_id=customer.id,
                            promotion_type="Fictional promo", window_start=date(2026, 1, 1), window_end=date(2026, 1, 31))
            db.add(case)
            cases.append(case)
        await db.flush()
        # USD / estimate / result: NB=(40,30,10), NR=(60,10,2).
        for i, (case, product, usd, est, sold) in enumerate([
            (cases[0], products[0], 10, 10, 8), (cases[0], products[1], 60, 10, 2),
            (cases[1], products[0], 30, 20, 2),
        ]):
            db.add(CporCaseLine(case_id=case.id, product_id=product.id, window_start=date(2026, 1, 1),
                               window_end=date(2026, 1, 31), srp=100, vat_rate=15, dealer_margin_pct=10,
                               estimate_qty=est, result_qty=sold, ttl_support_usd=usd, ttl_support=usd * 18))
        await db.flush()
        token = request_tenant_id.set(tenant)
        try:
            clear_query_cache()
            # Independent expected components; ratios recomputed over the filtered cohort.
            scenarios = [({}, (100, 40, 12)), ({"bu": " nb "}, (40, 30, 10)),
                         ({"customer_id": customers[0].id}, (70, 20, 10)),
                         ({"customer_id": customers[0].id, "bu": "NB"}, (10, 10, 8)),
                         ({"customer_id": -1}, (0, 0, 0)), ({"bu": "absent"}, (0, 0, 0))]
            for metric in ("support_spend", "delivery_rate", "support_cost_per_unit_sold"):
                for grain in load_catalog().metric_by_key(metric).allowed_grains:
                    # Period and cross-tab refusals stay legitimate, checked below.
                    if "period" in grain or len(grain) != 1:
                        continue
                    for filters, (usd, est, sold) in scenarios:
                        result = await execute_query(db, metric=metric, grains=list(grain), filters=filters,
                                                     tenant_id=tenant, skip_cache=True)
                        assert result.ok, (metric, grain, filters, result.message)
                        expected = usd if metric == "support_spend" else (
                            sold / est if metric == "delivery_rate" and est else
                            usd / sold if metric == "support_cost_per_unit_sold" and sold else None)
                        if expected is None:
                            assert result.value is None
                        else:
                            assert result.value == pytest.approx(expected), (metric, grain, filters)
                        assert result.scorecard["support_usd"] == usd, (grain, filters, result.as_dict())
                        assert sum(row["support_usd"] for row in result.rows) == usd
                        if not est:
                            assert result.rows == []
                        v2filters = {"customer" if key == "customer_id" else key: value for key, value in filters.items()}
                        v2 = await run_v2(db, V2Request(measures=[metric], breakdowns=list(grain), filters=v2filters), tenant_id=tenant)
                        assert v2["ok"], v2
                        assert v2["totals"][metric] == result.value
                        assert [r[metric] for r in v2["rows"]] == [r["value"] for r in result.rows]

            # A previously warmed unfiltered key must not bypass refusal checks.
            await execute_query(db, metric="support_spend", grains=["customer"], tenant_id=tenant)
            for invalid in [{"bu": ""}, {"bu": None}, {"unexpected": None}, {"unexpected": ""},
                            {"product_line": ""}, {"customer_id": ""}]:
                answer = await execute_query(db, metric="support_spend", grains=["customer"], filters=invalid, tenant_id=tenant)
                assert answer.status == "refused", (invalid, answer.as_dict())

            # Actual export / delivery owning functions consume the same filtered result.
            filters = {"bu": "NB"}
            result = await execute_query(db, metric="support_spend", grains=["customer"], filters=filters, tenant_id=tenant)
            content = build_xlsx_bytes(title="E1", metric_key="support_spend", grains=["customer"], value=result.value,
                                       rows=result.rows, data_vintage=result.data_vintage)
            book = load_workbook(BytesIO(content))
            assert book["Cover"]["A4"].value == "Value: 40.0"
            report = SavedReport(tenant_id=tenant, name="E1 fictional", metric_key="support_spend", grains=["customer"], filters=filters)
            db.add(report)
            await db.flush()
            monkeypatch.setattr(db, "commit", db.flush)
            delivered = await deliver_saved_report(db, report=report, user=None)
            assert delivered.status == "delivered" and delivered.value_preview == "40.0"
            definition = {"version": 2, "query": {"measures": ["support_spend"], "breakdowns": ["customer"], "filters": filters}, "display": {"visual": "table"}}
            payload = await _v2_payload(db, definition, tenant_id=tenant, skip_cache=True)
            assert payload["ok"] and payload["value"] == 40.0

            # Product scope is an intersection with the filter, never a substitute for it.
            for scope, expected in [(frozenset({"NR"}), 0), (frozenset(), 0), (frozenset({"*"}), 40)]:
                read_token, role_token = request_product_lines.set(scope), request_role.set("viewer")
                try:
                    from app.core.product_line_scope import apply_product_line_guc
                    await apply_product_line_guc(db)
                    # Force fresh ORM objects; preloaded relationships must not mask RLS.
                    db.expunge_all()
                    scoped = await execute_query(db, metric="support_spend", grains=["customer"], filters=filters, tenant_id=tenant, skip_cache=True)
                    assert scoped.value == expected
                finally:
                    request_product_lines.reset(read_token)
                    request_role.reset(role_token)
        finally:
            request_tenant_id.reset(token)
            await db.rollback()
            clear_query_cache()


@pytest.mark.anyio
@pytest.mark.parametrize("filters", [{"product_id": 1}, {"promo_type": "x"}, {"customer_id": "garbage"},
                                       {"customer_id": [1]}, {"bu": ["NB"]}, {"bu": "NB", "product_line": "NR"}])
async def test_cpor_refuses_unsupported_or_ambiguous_filters_without_io(filters):
    from app.query.handlers.a2_cpor import handle_a2
    from app.query.types import QueryRequest
    result = await handle_a2(None, QueryRequest(metric="support_spend", grains=["customer"], filters=filters),
                             metric_key="support_spend", explain_only=True)
    assert result.status == "refused", result


@pytest.mark.anyio
async def test_cpor_retains_all_time_and_crosstab_refusals():
    from app.query.handlers.a2_cpor import handle_a2
    from app.query.types import QueryRequest
    for grains, filters, phrase in [(["customer"], {"period_from": "2026-01-01"}, "no period filter"),
                                    (["customer", "bu"], {}, "cross-tab")]:
        result = await handle_a2(None, QueryRequest(metric="support_spend", grains=grains, filters=filters),
                                 metric_key="support_spend", explain_only=True)
        assert result.status == "not_implemented" and phrase in result.message


# Reuse the existing bearer harness: real auth and forced RLS; all accounts/data
# stay inside its outer rollback. This is API integration, not browser smoke.
from tests.test_dashboard_v2_authenticated import signed_in  # noqa: E402, F401


@pytest.mark.anyio
async def test_cpor_authenticated_query_export_and_saved_delivery(signed_in):
    from sqlalchemy import select
    from app.api.deps import get_db
    from app.main import app

    fixture = signed_in
    async for db in app.dependency_overrides[get_db]():
        for i, product in enumerate(fixture.products):
            customer = (await db.scalars(select(DimCustomer).where(DimCustomer.tenant_id == product.tenant_id))).first()
            case = CporCase(tenant_id=product.tenant_id, case_code=f"E1-http-{uuid4().hex[:15]}",
                            customer_id=customer.id, promotion_type="Fictional HTTP promo",
                            window_start=date(2026, 1, 1), window_end=date(2026, 1, 31))
            db.add(case)
            await db.flush()
            db.add(CporCaseLine(case_id=case.id, product_id=product.id, window_start=case.window_start,
                               window_end=case.window_end, srp=100, vat_rate=15, dealer_margin_pct=10,
                               estimate_qty=10, result_qty=5, ttl_support_usd=[10, 60, 1000][i], ttl_support=180))
        await db.commit()  # Existing harness savepoint; the outer transaction owns rollback.
    body = {"metric": "support_spend", "grains": ["customer"], "filters": {"bu": fixture.line_a}}
    v2 = {"measures": ["support_spend"], "breakdowns": ["customer"], "filters": body["filters"]}
    for key, total in [("admin", 10), ("viewer", 10), ("none", 0), ("other", 1000), ("viewer", 10)]:
        headers = fixture.headers[key]
        old = await fixture.client.post("/api/v1/query/execute", json=body, headers=headers)
        new = await fixture.client.post("/api/v1/query/v2/execute", json=v2, headers=headers)
        assert old.status_code == new.status_code == 200
        assert old.json()["ok"] and old.json()["value"] == total
        assert new.json()["ok"] and new.json()["totals"]["support_spend"] == total
        if key == "none":
            assert old.json()["rows"] == new.json()["rows"] == []
    headers = fixture.headers["admin"]
    for definition in [None, {"v": 1, "query": {**v2, "limit": 1}}]:
        report_body = {"name": "E1 HTTP fictional", **body}
        if definition:
            report_body["definition_v2"] = definition
        response = await fixture.client.post("/api/v1/saved-reports", json=report_body, headers=headers)
        assert response.status_code == 200, response.text
        report_id = response.json()["id"]
        export = await fixture.client.get(f"/api/v1/reports/saved/{report_id}/export?format=xlsx", headers=headers)
        assert export.status_code == 200
        assert load_workbook(BytesIO(export.content))["Cover"]["A4"].value == "Value: 10.0"
        delivered = await fixture.client.post(f"/api/v1/reports/saved/{report_id}/deliver", json={"format": "xlsx"}, headers=headers)
        assert delivered.status_code == 200, delivered.text
        assert delivered.json()["status"] == "delivered" and delivered.json()["value_preview"] == "10.0"
    for fmt in ("xlsx", "pdf"):
        response = await fixture.client.post("/api/v1/reports/export", json={**body, "format": fmt}, headers=headers)
        assert response.status_code == 200
        if fmt == "xlsx":
            assert load_workbook(BytesIO(response.content))["Cover"]["A4"].value == "Value: 10.0"
        else:
            assert b"Value: 10.0" in response.content
