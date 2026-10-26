"""E2: untrusted workbook text stays literal through common and HTTP exporters."""
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.services.report_export import build_xlsx_bytes
from tests.test_dashboard_v2_authenticated import signed_in  # noqa: F401


LABELS = ["=1+1", '=HYPERLINK("https://example.invalid/", "fictional")', "#N/A",
          "+1+1", "-1+1", "@SUM(1,1)", "\t=1+1", "Ordinary text"]


@pytest.mark.parametrize("label", LABELS)
def test_workbook_preserves_literal_titles_headers_labels_and_scalar_strings(label):
    for rows in (None, [{label: label, "amount": -123.45, "units": 17, "flag": True, "missing": None}]):
        content = build_xlsx_bytes(title=label, metric_key="fictional", grains=["customer"],
                                   value=label, rows=rows, data_vintage={})
        book = load_workbook(BytesIO(content), data_only=False)
        assert book["Cover"]["A1"].value == label
        assert book["Cover"]["A1"].data_type == "s"
        if rows is None:
            assert book["Data"]["A2"].value == label
            assert book["Data"]["A2"].data_type == "s"
        else:
            assert book["Data"]["A1"].value == book["Data"]["A2"].value == label
            assert book["Data"]["A1"].data_type == book["Data"]["A2"].data_type == "s"
            assert book["Data"]["B2"].value == -123.45 and book["Data"]["B2"].data_type == "n"
            assert book["Data"]["C2"].value == 17 and book["Data"]["C2"].data_type == "n"
            assert book["Data"]["D2"].value is True and book["Data"]["D2"].data_type == "b"
            assert book["Data"]["E2"].value is None
        assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
        with ZipFile(BytesIO(content)) as archive:
            assert not any("externalLinks" in name for name in archive.namelist())
            for name in ["xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"]:
                assert b"<f>" not in archive.read(name)
        # A reader expecting cached values sees the literal, not an empty formula result.
        assert load_workbook(BytesIO(content), data_only=True)["Data"]["A2"].value == label


@pytest.mark.parametrize("number", [0, -12, 13.625, False, None])
def test_scalar_numbers_booleans_and_missing_values_keep_their_types(number):
    content = build_xlsx_bytes(title="Fictional", metric_key="fictional", grains=[], value=number, rows=None, data_vintage={})
    value = load_workbook(BytesIO(content))["Data"]["A2"].value
    assert value == number
    assert type(value) is type(number)


@pytest.mark.anyio
async def test_authenticated_ad_hoc_legacy_and_model_saved_exports_keep_labels_literal(signed_in):
    from sqlalchemy import select
    from app.api.deps import get_db
    from app.main import app
    from app.models.dimensions import DimCustomer

    fixture = signed_in
    label = '=HYPERLINK("https://example.invalid/", "fictional")'
    async for db in app.dependency_overrides[get_db]():
        customer = (await db.scalars(select(DimCustomer).where(DimCustomer.tenant_id == fixture.products[0].tenant_id))).first()
        customer.name = label
        await db.commit()  # Existing harness savepoint; outer transaction rolls back.
    headers = fixture.headers["admin"]
    query = {"metric": "sellout_units", "grains": ["customer"], "filters": {"product_id": fixture.products[0].id}}
    content = []
    response = await fixture.client.post("/api/v1/reports/export", json={**query, "title": "=1+1", "format": "xlsx"}, headers=headers)
    assert response.status_code == 200, response.text
    content.append(response.content)
    for definition in [None, {"v": 1, "query": {"measures": ["dsi_sellout.units", "dsi_sellout.lines"],
                                              "breakdowns": ["customer"], "filters": {"product": [fixture.products[0].id]}}}]:
        body = {"name": "=1+1", **query}
        if definition is not None:
            body["definition_v2"] = definition
        saved = await fixture.client.post("/api/v1/saved-reports", json=body, headers=headers)
        assert saved.status_code == 200, saved.text
        if definition is not None:
            assert saved.json()["definition_v2"]["v1_equivalent"] is False
        response = await fixture.client.get(f"/api/v1/reports/saved/{saved.json()['id']}/export?format=xlsx", headers=headers)
        assert response.status_code == 200, response.text
        content.append(response.content)
    for exported in content:
        book = load_workbook(BytesIO(exported))
        assert book["Cover"]["A1"].value == "=1+1" and book["Cover"]["A1"].data_type == "s"
        cells = [cell for sheet in book for row in sheet for cell in row]
        labels = [cell for cell in cells if cell.value == label]
        assert labels and all(cell.data_type == "s" for cell in labels)
        assert all(cell.data_type != "f" for cell in cells)
        assert "913" in book["Cover"]["A4"].value
